"""Reading an event's running order into rows of sessions.

Spreadsheets and CSV files with recognisable column headings (Room, Session,
Start, End, Speaker...) are read directly. Anything else (PDF, Word, a
spreadsheet without headings) goes to a local AI model through Ollama, if one
is set up in Admin → Presenter; nothing leaves the venue network.
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, time, timedelta

import httpx

FIELDS = ("room_name", "title", "starts_at", "ends_at", "presenter_name", "presenter_email", "presenter_phone")
HEADINGS = {
    "room_name": ("room", "room name", "venue", "location", "stage", "track", "hall"),
    "title": ("title", "session", "session title", "session name", "talk", "topic", "item", "agenda item"),
    "date": ("date", "day"),
    "start": ("start", "starts", "start time", "starts at", "time", "from", "begin"),
    "end": ("end", "ends", "end time", "ends at", "to", "finish", "until"),
    "duration": ("duration", "length", "mins", "minutes", "duration (mins)", "duration (min)"),
    "presenter_name": ("presenter", "speaker", "speakers", "presenter name", "speaker name", "presented by", "name"),
    "presenter_email": ("email", "e-mail", "presenter email", "speaker email"),
    "presenter_phone": ("phone", "mobile", "telephone", "presenter phone", "speaker phone"),
}


class ScheduleError(Exception):
    pass


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9()]+", " ", str(h or "").lower()).strip()


def _table(filename: str, data: bytes) -> list[list] | None:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("csv", "txt", "tsv"):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if ext == "tsv" else csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        for ws in wb.worksheets:  # the first sheet with a heading row wins
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if _header(rows):
                return rows
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)] if wb.worksheets else []
    return None


def _header(rows: list[list]) -> tuple[int, dict] | None:
    """Find the heading row in the first 15 rows: it needs a title column and
    at least one of room or start."""
    for i, row in enumerate(rows[:15]):
        cols = {}
        for j, cell in enumerate(row):
            h = _norm(cell)
            for field, names in HEADINGS.items():
                if field not in cols and h in names:
                    cols[field] = j
                    break
        if "title" in cols and ("room_name" in cols or "start" in cols):
            return i, cols
    return None


def _date(v, default: date | None) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return default
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y", "%a %d %b %Y", "%A %d %B %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    m = re.match(r"^day\s*(\d+)$", s, re.I)
    if m and default:
        return default + timedelta(days=int(m.group(1)) - 1)
    return default


def _time(v) -> time | None:
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    if isinstance(v, (int, float)) and 0 <= v < 1:  # Excel day fraction
        mins = round(v * 24 * 60)
        return time(mins // 60 % 24, mins % 60)
    s = str(v or "").strip().lower().replace(".", ":")
    m = re.match(r"^(\d{1,2})(?::(\d{2}))?(?::\d{2})?\s*(am|pm)?$", s)
    if not m:
        m = re.search(r"(\d{1,2}):(\d{2})\s*(am|pm)?", s)
        if not m:
            return None
    h, mi, ap = int(m.group(1)), int(m.group(2) or 0), m.group(3)
    if ap == "pm" and h < 12:
        h += 12
    if ap == "am" and h == 12:
        h = 0
    return time(h, mi) if h < 24 and mi < 60 else None


def _iso(d: date | None, t: time | None) -> str:
    if not t:
        return ""
    return datetime.combine(d, t).isoformat(timespec="minutes") if d else t.strftime("%H:%M")


def from_table(rows: list[list], event_start: str = "") -> list[dict] | None:
    found = _header(rows)
    if not found:
        return None
    head, cols = found
    first_day = _date(event_start[:10], None) if event_start else None
    day = first_day
    get = lambda r, f: r[cols[f]] if f in cols and cols[f] < len(r) else None  # noqa: E731
    out = []
    for r in rows[head + 1:]:
        if not any(str(c or "").strip() for c in r):
            continue
        if "date" in cols:
            day = _date(get(r, "date"), first_day) or day
        start_v = get(r, "start")
        # "09:30 - 10:15" in one cell
        if isinstance(start_v, str) and re.search(r"(\d|[ap]m)\s*[-–]\s*\d", start_v, re.I) and "end" not in cols:
            a, b = re.split(r"\s*[-–]\s*", start_v, maxsplit=1)
            start_t, end_t = _time(a), _time(b)
        else:
            start_t, end_t = _time(start_v), _time(get(r, "end"))
        sd = start_v.date() if isinstance(start_v, datetime) else day
        if not end_t and start_t and "duration" in cols:
            try:
                end_t = (datetime.combine(date.today(), start_t) + timedelta(minutes=float(get(r, "duration")))).time()
            except (TypeError, ValueError):
                pass
        row = {
            "room_name": str(get(r, "room_name") or "").strip(),
            "title": str(get(r, "title") or "").strip(),
            "starts_at": _iso(sd, start_t),
            "ends_at": _iso(sd, end_t),
            "presenter_name": str(get(r, "presenter_name") or "").strip(),
            "presenter_email": str(get(r, "presenter_email") or "").strip(),
            "presenter_phone": str(get(r, "presenter_phone") or "").strip(),
        }
        if row["title"]:
            out.append(row)
    return out


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    table = _table(filename, data)
    if table is not None:
        return "\n".join(" | ".join(str(c) for c in r if c not in (None, "")) for r in table)
    raise ScheduleError(f'Can\'t read ".{ext}" files: upload a spreadsheet, CSV, PDF or Word file')


PROMPT = """You are extracting a structured running order from an event schedule for the event "{event}" (starts {starts}).

Rooms already set up: {rooms}. When a row's room matches one of these (allowing for small spelling or case differences), use the exact known name; otherwise use the room name as written.

Return ONLY a JSON array. Each element is an object with these keys:
  "room_name", "title", "starts_at" (ISO 8601 like "2026-10-02T09:30", resolve "Day 1" against the event start, "" if not stated),
  "ends_at" (ISO 8601 or ""), "presenter_name", "presenter_email", "presenter_phone" ("" when not given)

Document:
---
{text}
---

JSON array:"""


def with_ai(text: str, url: str, model: str, event: str, starts: str, rooms: list[str]) -> list[dict]:
    if not url:
        raise ScheduleError("This file needs the schedule AI, which isn't set up. Use a spreadsheet or CSV with column "
                            "headings (Room, Session, Start, End, Speaker), or set up Ollama in Admin → Presenter.")
    prompt = PROMPT.format(event=event or "this event", starts=starts or "unknown", rooms=", ".join(rooms) or "none yet",
                           text=text[:20000])
    try:
        r = httpx.post(f"{url.rstrip('/')}/api/generate", json={"model": model, "prompt": prompt, "format": "json", "stream": False},
                       timeout=240)
        r.raise_for_status()
    except httpx.HTTPError as e:
        raise ScheduleError(f"Couldn't reach the schedule AI at {url} ({e})") from None
    raw = (r.json().get("response") or "").strip()
    m = re.match(r"^```(?:json)?\s*(.*?)\s*```$", raw, re.S)
    try:
        rows = json.loads(m.group(1) if m else raw)
    except ValueError:
        raise ScheduleError("The schedule AI didn't return a readable list. Try again, or use a spreadsheet.") from None
    if isinstance(rows, dict):  # some models wrap the list
        rows = next((v for v in rows.values() if isinstance(v, list)), [])
    if not isinstance(rows, list):
        raise ScheduleError("The schedule AI didn't return a list of sessions")
    return [{f: str(r.get(f) or "").strip() for f in FIELDS} for r in rows if isinstance(r, dict) and r.get("title")]


def parse(filename: str, data: bytes, *, event: str = "", starts: str = "", rooms: list[str] | None = None,
          ai_url: str = "", ai_model: str = "") -> tuple[list[dict], str]:
    """Rows and how they were read ("table" or "ai")."""
    table = _table(filename, data)
    if table is not None:
        rows = from_table(table, starts)
        if rows is not None:
            return rows, "table"
    return with_ai(extract_text(filename, data), ai_url, ai_model, event, starts, rooms or []), "ai"
